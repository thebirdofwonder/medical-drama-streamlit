"""
論文PDF / URL から、冒頭文・参考文献・タイトル・台本・画像を作る前半工程。
APIキーは環境変数 / .env / 画面入力から読む（コードに直書きしない）。
"""

from __future__ import annotations

import html as html_lib
import io
import json
import os
import random
import re
import urllib.parse
from pathlib import Path
from typing import Any, Callable

import requests
from docx import Document
from PIL import Image, ImageDraw, ImageEnhance, ImageFilter

# app.py 側の共通関数は呼び出し時に渡す／遅延 import する


CLAUDE_MODEL_DEFAULT = "claude-sonnet-4-20250514"
TITLE_ACCENTS = [
    ("red", (140, 36, 36)),
    ("orange", (150, 78, 24)),
    ("yellow", (130, 110, 28)),
    ("green", (28, 100, 52)),
    ("blue", (28, 56, 120)),
    ("indigo", (48, 40, 110)),
    ("purple", (92, 36, 110)),
]


def load_dotenv_file(path: Path) -> None:
    """簡易 .env 読み込み（python-dotenv なし）。既存の環境変数は上書きしない。"""
    if not path.is_file():
        return
    try:
        for line in path.read_text(encoding="utf-8").splitlines():
            s = line.strip()
            if not s or s.startswith("#") or "=" not in s:
                continue
            key, val = s.split("=", 1)
            key = key.strip()
            val = val.strip().strip('"').strip("'")
            if key and key not in os.environ:
                os.environ[key] = val
    except Exception:
        pass


def get_anthropic_api_key(explicit: str = "") -> str:
    """Anthropic APIキーを安全に取得（直書きしない）。"""
    load_dotenv_file(Path(__file__).resolve().parent / ".env")
    for cand in (
        (explicit or "").strip(),
        (os.environ.get("ANTHROPIC_API_KEY") or "").strip(),
    ):
        if cand:
            return cand
    try:
        import streamlit as st

        v = str(st.session_state.get("anthropic_api_key") or "").strip()
        if v:
            return v
        secrets = getattr(st, "secrets", None)
        if secrets is not None:
            v = str(secrets.get("ANTHROPIC_API_KEY", "") or "").strip()
            if v:
                return v
    except Exception:
        pass
    return ""


def save_anthropic_api_key_to_dotenv(api_key: str) -> Path:
    """APIキーを .env に保存（.gitignore 済み）。"""
    root = Path(__file__).resolve().parent
    env_path = root / ".env"
    key = (api_key or "").strip()
    lines: list[str] = []
    if env_path.is_file():
        for line in env_path.read_text(encoding="utf-8").splitlines():
            if line.strip().startswith("ANTHROPIC_API_KEY="):
                continue
            lines.append(line)
    lines.append(f"ANTHROPIC_API_KEY={key}")
    env_path.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")
    os.environ["ANTHROPIC_API_KEY"] = key
    return env_path


def http_session_direct() -> requests.Session:
    session = requests.Session()
    session.trust_env = False
    session.proxies = {"http": None, "https": None}
    return session


def strip_code_fence(text: str) -> str:
    text = (text or "").strip()
    if text.startswith("```"):
        text = re.sub(r"^```(?:json|JSON)?\s*", "", text)
        text = re.sub(r"\s*```$", "", text)
    return text.strip()


def claude_messages(
    prompt: str,
    *,
    api_key: str,
    model: str = CLAUDE_MODEL_DEFAULT,
    max_tokens: int = 8000,
    system: str = "",
) -> str:
    """Anthropic Messages API を呼び、テキストを返す。"""
    key = (api_key or "").strip()
    if not key:
        raise RuntimeError(
            "Anthropic APIキーがありません。"
            " 左の設定でキーを入れるか、.env に ANTHROPIC_API_KEY=... を書いてください。"
        )
    body: dict[str, Any] = {
        "model": model or CLAUDE_MODEL_DEFAULT,
        "max_tokens": int(max_tokens),
        "messages": [{"role": "user", "content": prompt}],
    }
    if system.strip():
        body["system"] = system.strip()
    r = http_session_direct().post(
        "https://api.anthropic.com/v1/messages",
        headers={
            "x-api-key": key,
            "anthropic-version": "2023-06-01",
            "content-type": "application/json",
        },
        json=body,
        timeout=180,
    )
    if r.status_code >= 400:
        raise RuntimeError(
            f"Claude API エラー HTTP {r.status_code}: {r.text[:500]}"
        )
    data = r.json()
    parts = data.get("content") or []
    texts = [
        str(p.get("text") or "")
        for p in parts
        if isinstance(p, dict) and p.get("type") == "text"
    ]
    out = "\n".join(t for t in texts if t).strip()
    if not out:
        raise RuntimeError("Claude から空の応答が返りました。")
    return out


def fetch_paper_from_url(url: str, extract_pdf_bytes: Callable[[bytes], str]) -> str:
    """論文URLから本文テキストを取得（HTMLまたはPDF）。"""
    u = (url or "").strip()
    if not u:
        raise ValueError("URLが空です。")
    if not re.match(r"^https?://", u, re.I):
        u = "https://" + u
    r = http_session_direct().get(
        u,
        timeout=90,
        headers={"User-Agent": "Mozilla/5.0 (compatible; MedicalDramaApp/1.0)"},
        allow_redirects=True,
    )
    r.raise_for_status()
    ctype = (r.headers.get("Content-Type") or "").lower()
    if "pdf" in ctype or u.lower().endswith(".pdf") or r.content[:4] == b"%PDF":
        return extract_pdf_bytes(r.content)
    raw = r.content
    # 文字コード推定
    try:
        text_html = raw.decode(r.encoding or "utf-8", errors="replace")
    except Exception:
        text_html = raw.decode("utf-8", errors="replace")
    text_html = re.sub(r"(?is)<script[^>]*>.*?</script>", " ", text_html)
    text_html = re.sub(r"(?is)<style[^>]*>.*?</style>", " ", text_html)
    text_html = re.sub(r"(?is)<noscript[^>]*>.*?</noscript>", " ", text_html)
    text = re.sub(r"(?is)<[^>]+>", " ", text_html)
    text = html_lib.unescape(text)
    text = re.sub(r"[ \t\u3000]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def build_paper_generation_prompt(paper_text: str) -> str:
    """論文本文から opening / reference / title / script を作る指示。"""
    body = (paper_text or "").strip()
    if len(body) > 60000:
        body = body[:60000] + "\n…(以下省略)"
    return f"""
あなたは医学教育用の医療ドラマ台本作家です。次の症例報告・論文をもとに、指定の成果物を日本語で作ってください。

【論文本文】
{body}

【出力形式】
必ず JSON オブジェクトだけを返す（説明文やコードフェンスは付けない）。キーは次のとおり:
{{
  "opening": "冒頭4〜5文を1段落にした文章",
  "reference": "YouTube説明欄用の論文情報＋著作権情報",
  "titles": ["案1", "案2", "案3"],
  "best_title": "一番の推奨タイトル（titlesのいずれか）",
  "script": "完成台本（改行あり）",
  "image_prompts": {{
    "title_card": "タイトル画用の英語プロンプト（文字なし・暗め・知的）",
    "backgrounds": [
      {{"hint": "手術室の風景", "prompt": "English scene prompt without text"}}
    ]
  }}
}}

【opening の条件】
- 4〜5文でストーリーの出だしだけ。
- 本当の診断名は明示しない。診断につながるヒントも示唆しない。
- 改行なし。1段落の文章。

【reference の条件】
- YouTubeの説明欄にそのまま貼れる形式。
- 論文の書誌情報（著者・タイトル・雑誌・年・DOI/URL等、分かる範囲）と著作権・引用の注意を含める。
- 医療ドラマ化であること、教育目的のフィクションであることも短く書く。

【titles / best_title の条件】
次の例と同じ長さ・体裁・要領のキャッチーなタイトルを複数案出し、best_title に最良案を入れる。
例:
•  【症例報告／医療ドラマ】◯◯歳女性の反復する激しい［症状］…前医の「右卵巣嚢胞」診断と画像不一致の違和感｜卵巣嚢胞の向こう側
•  【症例報告／医療ドラマ】◯◯歳女性の［所見］…穿刺目的で紹介された患者の正常卵巣が告げる予期せぬ真相｜卵巣嚢胞の向こう側
（この論文内容に合わせて具体化する。シリーズ名部分は論文の主題に合うものにする）

【script の条件】
① プロット全体の構成や意外性のある展開になるよう brush up。ナレーターが読み上げる本文以外の余計な文・記号・見出しは入れない。［体言止め＋読点］を多用しない。列挙は読点で並べる。短く名詞を句点で並列しない。
② この分野の専門医が読んで違和感・不適切用語・論理破綻があれば修正。検査値の単位は原文のまま変更しない。
③ 最後の「解説」は長くしすぎず、全体のおよそ10%にまとめる。
④ 数段落ごと、風景が変わると良い箇所の文頭にだけ 〈風景描写〉 を入れる。例: 〈手術室の風景〉 〈CTを撮影する部屋の風景〉 〈心電図画面〜波形はVTを表示〉 〈ERに担ぎ込まれた急患を取り囲む医師、看護師たち〉。毎段落必須ではない。
⑤ 台本はナレーション本文のみ（ト書きや「ナレーター:」などのラベルは付けない）。

【image_prompts】
- title_card: 文字なし。暗めだが少し明るめ。知的な医師・医療の雰囲気。おどろおどろしくしない。症例内容を反映。
- backgrounds: script 中のすべての 〈〉 ヒントについて、hint（中身）と英語 prompt を列挙。文字・ロゴ・字幕を入れない。
""".strip()


def parse_generation_json(raw: str) -> dict[str, Any]:
    text = strip_code_fence(raw)
    try:
        data = json.loads(text, strict=False)
    except json.JSONDecodeError:
        m = re.search(r"\{[\s\S]*\}", text)
        if not m:
            raise RuntimeError("Claude の応答を JSON として読めませんでした。")
        data = json.loads(m.group(0), strict=False)
    if not isinstance(data, dict):
        raise RuntimeError("JSON の形が不正です。")
    return data


def save_text_docx(path: Path, text: str) -> Path:
    """テキストを docx で保存（段落分割）。"""
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    content = (text or "").replace("\r\n", "\n").strip()
    if not content:
        doc.add_paragraph("")
    else:
        for para in content.split("\n"):
            doc.add_paragraph(para)
    doc.save(path)
    return path


def extract_angle_hints(script: str) -> list[str]:
    """台本中の 〈〉 系ヒントの中身を順に返す。"""
    open_chars = "＜<〈‹《〈⟨❮"
    close_chars = "＞>〉›》〉⟩❯"
    pat = re.compile(
        rf"[{re.escape(open_chars)}]"
        rf"[^{re.escape(open_chars + close_chars)}]{{1,200}}"
        rf"[{re.escape(close_chars)}]"
    )
    hints: list[str] = []
    for m in pat.finditer(script or ""):
        inner = (m.group(0) or "")[1:-1].strip()
        inner = re.sub(r"\s+", " ", inner)
        if inner and inner not in hints:
            hints.append(inner)
    return hints


def _pillow_tinted_medical_card(
    size: tuple[int, int],
    accent: tuple[int, int, int],
    *,
    brighter: bool = True,
) -> Image.Image:
    """API画像が取れないときの、暗め＋色味付きの知的な医療イメージ。"""
    w, h = size
    ar, ag, ab = accent
    img = Image.new("RGB", (w, h), (10, 12, 18))
    draw = ImageDraw.Draw(img)
    for y in range(h):
        t = y / max(h - 1, 1)
        r = int(12 + ar * (0.45 + 0.25 * t))
        g = int(14 + ag * (0.40 + 0.22 * t))
        b = int(20 + ab * (0.50 + 0.20 * t))
        draw.line([(0, y), (w, y)], fill=(min(220, r), min(220, g), min(220, b)))
    # 抽象的な円・弧（知的・医療っぽさ。文字なし）
    overlay = Image.new("RGBA", (w, h), (0, 0, 0, 0))
    od = ImageDraw.Draw(overlay)
    od.ellipse(
        [int(w * 0.55), int(h * -0.1), int(w * 1.05), int(h * 0.55)],
        outline=(255, 255, 255, 35),
        width=3,
    )
    od.ellipse(
        [int(w * -0.15), int(h * 0.45), int(w * 0.45), int(h * 1.15)],
        outline=(255, 255, 255, 28),
        width=2,
    )
    # 簡易ECG風ライン
    y0 = int(h * 0.62)
    pts = []
    x = int(w * 0.08)
    while x < int(w * 0.92):
        pts.extend(
            [
                (x, y0),
                (x + 18, y0),
                (x + 26, y0 - 16),
                (x + 34, y0 + 55),
                (x + 42, y0 - 70),
                (x + 50, y0 + 18),
                (x + 58, y0),
                (x + 96, y0),
            ]
        )
        x += 110
    if len(pts) >= 2:
        od.line(pts, fill=(220, 240, 255, 70), width=3)
    img = Image.alpha_composite(img.convert("RGBA"), overlay).convert("RGB")
    if brighter:
        img = ImageEnhance.Brightness(img).enhance(1.12)
        img = ImageEnhance.Contrast(img).enhance(1.05)
    img = img.filter(ImageFilter.SMOOTH)
    return img


def download_pollinations_image(
    prompt: str,
    size: tuple[int, int] = (1920, 1080),
) -> Image.Image | None:
    """無料の画像生成エンドポイントを試し、失敗時は None。"""
    w, h = size
    q = urllib.parse.quote((prompt or "medical drama scene").strip()[:900])
    url = (
        f"https://image.pollinations.ai/prompt/{q}"
        f"?width={w}&height={h}&nologo=true&enhance=true"
    )
    try:
        r = http_session_direct().get(url, timeout=120, allow_redirects=True)
        if r.status_code != 200 or len(r.content) < 3000:
            return None
        ctype = (r.headers.get("Content-Type") or "").lower()
        if "html" in ctype:
            return None
        with Image.open(io.BytesIO(r.content)) as im:
            im.load()
            return im.convert("RGB").resize(size, Image.Resampling.LANCZOS)
    except Exception:
        return None


def generate_scene_image(
    prompt: str,
    *,
    size: tuple[int, int] = (1920, 1080),
    accent: tuple[int, int, int] | None = None,
) -> Image.Image:
    """情景画像を生成（外部API → 失敗時は色味付きPilow）。"""
    img = download_pollinations_image(prompt, size=size)
    if img is not None:
        # 白文字が載せやすいよう、やや暗く整える
        img = ImageEnhance.Brightness(img).enhance(0.82)
        img = ImageEnhance.Color(img).enhance(0.95)
        return img
    name, color = random.choice(TITLE_ACCENTS)
    _ = name
    return _pillow_tinted_medical_card(size, accent or color, brighter=True)


def generate_title_card_image(
    prompt: str,
    *,
    size: tuple[int, int] = (1920, 1080),
) -> Image.Image:
    """YouTube用タイトル画（文字なし）。色味はランダムに濃い暖色〜寒色。"""
    name, accent = random.choice(TITLE_ACCENTS)
    full_prompt = (
        f"{prompt}. Cinematic medical drama mood, intelligent physician atmosphere, "
        f"no text, no letters, no logo, slightly bright but dark enough for white title overlay, "
        f"dominant deep {name} color grading, widescreen 16:9"
    )
    img = download_pollinations_image(full_prompt, size=size)
    if img is None:
        img = _pillow_tinted_medical_card(size, accent, brighter=True)
    else:
        # 選んだ色味を薄く重ねる
        tint = Image.new("RGB", size, accent)
        img = Image.blend(img.resize(size, Image.Resampling.LANCZOS), tint, 0.22)
        img = ImageEnhance.Brightness(img).enhance(0.88)
    return img


def run_paper_to_drama_pipeline(
    paper_text: str,
    *,
    api_key: str,
    desktop_dir: Path,
    custom_bg_dir: Path,
    video_size: tuple[int, int] = (1920, 1080),
    progress: Callable[[str], None] | None = None,
) -> dict[str, Any]:
    """
    論文テキストから成果物を一括作成し、デスクトップ／背景フォルダへ保存する。
    戻り値: opening, reference, best_title, script, paths, hints
    """

    def _p(msg: str) -> None:
        if progress:
            progress(msg)

    _p("Claude で台本・タイトル等を作成中…")
    raw = claude_messages(
        build_paper_generation_prompt(paper_text),
        api_key=api_key,
        system=(
            "You are a careful medical drama scriptwriter for educational YouTube videos. "
            "Return valid JSON only."
        ),
        max_tokens=10000,
    )
    data = parse_generation_json(raw)
    opening = str(data.get("opening") or "").strip()
    reference = str(data.get("reference") or "").strip()
    best_title = str(data.get("best_title") or "").strip()
    titles = data.get("titles") or []
    if not best_title and isinstance(titles, list) and titles:
        best_title = str(titles[0]).strip()
    script = str(data.get("script") or "").strip()
    if not script:
        raise RuntimeError("台本が空でした。もう一度生成してください。")
    if not opening:
        opening = script.split("。")[0] + "。" if "。" in script else script[:120]
    if not reference:
        reference = "（参考文献情報を取得できませんでした）"
    if not best_title:
        best_title = "【症例報告／医療ドラマ】診断の向こう側"

    desktop_dir.mkdir(parents=True, exist_ok=True)
    custom_bg_dir.mkdir(parents=True, exist_ok=True)

    _p("デスクトップへ Word を保存中…")
    paths = {
        "opening": save_text_docx(desktop_dir / "opening.docx", opening),
        "reference": save_text_docx(desktop_dir / "reference.docx", reference),
        "title": save_text_docx(desktop_dir / "title.docx", best_title),
        "script": save_text_docx(desktop_dir / "台本.docx", script),
    }

    image_prompts = data.get("image_prompts") or {}
    if not isinstance(image_prompts, dict):
        image_prompts = {}
    title_prompt = str(image_prompts.get("title_card") or "").strip()
    if not title_prompt:
        title_prompt = (
            "intellectual physician in modern hospital, cinematic lighting, no text"
        )

    _p("タイトル画を作成中…")
    title_img = generate_title_card_image(title_prompt, size=video_size)
    title_path = desktop_dir / "タイトル画.png"
    title_img.save(title_path, format="PNG")
    paths["title_image"] = title_path

    # 〈〉 ヒントごとに背景を作成
    bg_prompt_map: dict[str, str] = {}
    bg_list = image_prompts.get("backgrounds") or []
    if isinstance(bg_list, list):
        for item in bg_list:
            if not isinstance(item, dict):
                continue
            hint = str(item.get("hint") or "").strip()
            prompt = str(item.get("prompt") or "").strip()
            if hint:
                bg_prompt_map[hint] = prompt or hint

    hints = extract_angle_hints(script)
    for h in hints:
        bg_prompt_map.setdefault(
            h,
            f"medical drama scene: {h}, photorealistic hospital atmosphere, no text, no letters",
        )

    bg_paths: dict[str, Path] = {}
    for i, hint in enumerate(hints):
        _p(f"背景静止画 {i + 1}/{len(hints)}: {hint}")
        prompt = bg_prompt_map.get(hint) or hint
        img = generate_scene_image(
            f"{prompt}. No text, no watermark, widescreen medical still",
            size=video_size,
        )
        safe_name = re.sub(r'[\\/:*?"<>|\n\r]+', "", hint).strip() or f"bg_{i+1}"
        out = custom_bg_dir / f"{safe_name}.png"
        img.save(out, format="PNG")
        bg_paths[hint] = out

    return {
        "opening": opening,
        "reference": reference,
        "best_title": best_title,
        "titles": titles if isinstance(titles, list) else [],
        "script": script,
        "paths": paths,
        "bg_paths": bg_paths,
        "hints": hints,
    }
